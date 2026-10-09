# -*- coding: utf-8 -*-
"""Parse brand.xlsx → storage/brand_import.json"""
import json
import sys
from pathlib import Path

import openpyxl

root = Path(__file__).resolve().parent.parent
xlsx = root / (sys.argv[1] if len(sys.argv) > 1 else "brand.xlsx")
out = root / "storage" / "brand_import.json"

wb = openpyxl.load_workbook(xlsx, read_only=True)
data = []

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    brands = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 2:
            continue
        cells = list(row) + [None, None, None]
        ar = cells[1]
        en = cells[2]
        if ar is None or en is None:
            continue
        ar = str(ar).strip()
        en = str(en).strip()
        if not ar or not en:
            continue
        brands.append({"ar": ar, "en": en})
    data.append({"category_ar": sheet_name.strip(), "brands": brands})

out.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
print("Wrote", len(data), "categories,", sum(len(s["brands"]) for s in data), "brands to", out)
